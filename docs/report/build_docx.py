# -*- coding: utf-8 -*-
"""
통합 운영 모니터링 — 경영진/동료 대상 성과보고서 (DOCX)
발표자료(PPTX)와 동일 스토리라인. 정성 위주 · Cursor 효율성 중심.
"""
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

KOR = 'Malgun Gothic'
NAVY = RGBColor(0x0F, 0x1F, 0x3D)
BLUE = RGBColor(0x1E, 0x4D, 0x8C)
SLATE = RGBColor(0x5B, 0x67, 0x7A)


def _kor(run, size=None, bold=None, color=None):
    run.font.name = KOR
    run.element.rPr.rFonts.set(qn('w:eastAsia'), KOR)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def base(doc):
    st = doc.styles['Normal']
    st.font.name = KOR
    st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), KOR)
    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)


def h(doc, text, level=1):
    p = doc.add_heading(level=level)
    r = p.add_run(text)
    _kor(r, color=NAVY)
    return p


def b(doc, text, level=0):
    style = 'List Bullet' if level == 0 else 'List Bullet 2'
    p = doc.add_paragraph(style=style)
    r = p.add_run(text)
    _kor(r)
    return p


def para(doc, text, size=11, bold=False, color=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    _kor(r, size=size, bold=bold, color=color)
    return p


doc = Document()
base(doc)

# Title
t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
tr = t.add_run('AI 활용 개발 성과 보고서')
_kor(tr, 20, True, NAVY)

s = doc.add_paragraph()
s.alignment = WD_ALIGN_PARAGRAPH.CENTER
sr = s.add_run(
    'Cursor 기반 AI 페어프로그래밍으로 구축한 통합 운영 모니터링\n'
    '대상: 상무님 포함 사내 동료  ·  발표 기준 약 10분'
)
_kor(sr, 11, False, SLATE)
doc.add_paragraph()

h(doc, '1. 보고 목적', 1)
para(doc,
     '본 보고는 운영 점검 통합 프로그램의 기능 소개보다, '
     '개발 전 과정에 AI 도구(Cursor)와 표준 규칙(.md)을 적용하여 '
     '얻은 생산성·품질 효과를 중심으로 정리한다.')

h(doc, '2. 문제 인식', 1)
para(doc,
     '운영자는 매일 SM37(배치), ST22(런타임 덤프), SXI_MONITOR(인터페이스)를 '
     '개별 트랜잭션으로 확인해야 한다. 점검 채널이 분산되어 화면 전환이 잦고, '
     '누락·지연 리스크와 표준화된 점검 뷰의 부재가 상존한다.')

h(doc, '3. 솔루션 요약', 1)
b(doc, '단일 트랜잭션 통합 대시보드: 상단 요약 · 중간 Top-N 차트 · 하단 상세 ALV')
b(doc, '더블클릭 시 표준 상세 화면(잡 로그/덤프/메시지)으로 드릴다운')
b(doc, '읽기 전용 설계: 데이터 변경·COMMIT·LOCK 금지, 표준 테이블/FM만 사용')
b(doc, '영역별 권한 격리: 일부 영역 권한 부재 시에도 나머지 영역은 정상 동작')

h(doc, '4. AI(Cursor) 적용 방식', 1)
b(doc, '설계: 설계서 초안·개정 및 결정/미결 이력 관리')
b(doc, '분석: STAUTHTRACE 엑셀 분석으로 필요 권한 객체 도출')
b(doc, '조사: SAP 표준 FM·화면 사용법의 즉시 확인')
b(doc, '구현: 코드 생성 → 오류 진단 → 수정 루프 가속')
b(doc, '검증: 리뷰·표준 준수 점검, 문서/PR 품질 균일화')
para(doc,
     '역할 분담: AI는 조사·초안·반복 수정을 담당하고, '
     '개발자는 요구 확정·설계 의사결정·결과 검증·운영 책임을 담당한다.',
     bold=False)

h(doc, '5. Rules / Skills(.md) 기반 표준화', 1)
para(doc,
     '개발 표준과 작업 절차를 저장소의 .md 파일로 명문화하여, '
     '세션·담당자가 바뀌어도 동일 기준이 자동 적용되도록 하였다.')
h(doc, '5.1 Rules (상시 자동 적용)', 2)
b(doc, '.cursor/rules/abap-project-conventions')
b(doc, '읽기 전용 불변식, 표준 객체만 사용, 기간 제한 SELECT, 인터페이스 기반 OO', 1)
h(doc, '5.2 Skills (작업 유형별 플레이북)', 2)
b(doc, '.cursor/skills — 읽기전용 모니터링, 클린 OO, 성능, 코드리뷰, 디버깅, 설계서/커밋/PR')
h(doc, '5.3 실무 효과', 2)
b(doc, '표준 설명 반복 제거, 재작업·휴먼에러 감소, 리뷰/문서 품질 균일화, 온보딩 가속')
b(doc, '규칙 자체를 버전관리하여 팀 단위로 개선·재사용')

h(doc, '6. 효율성 · 성과 (정성 · 추정)', 1)
h(doc, '개발 생산성', 2)
b(doc, '설계–구현–검토를 단일 맥락으로 유지, 요구 변경의 즉시 반영')
h(doc, '리서치 부담', 2)
b(doc, '표준 FM/권한/화면 조사를 대화로 해결하여 탐색 시간 절감')
h(doc, '품질 · 표준', 2)
b(doc, 'Rules/Skills로 사내 표준 일관 적용, 재작업·품질 편차 축소')
para(doc,
     '효과 규모는 정성 추정이며, 향후 동일 난이도 과제에 대해 '
     '개발 소요시간 전/후 비교로 수치화할 예정이다.',
     color=SLATE)

h(doc, '7. 기대 효과 및 향후 계획', 1)
b(doc, '기대: 운영 점검 시간 단축·누락 방지, 운영 안전성, AI 적용 가능성 확인, 표준 자산화')
b(doc, '향후: 정식 오브젝트화, 정량 실측, 유사 과제 확대, 팀 AI 활용 가이드라인 정리')

h(doc, '8. 핵심 메시지', 1)
para(doc,
     '운영 점검을 통합하는 동시에, AI와 표준(.md)을 결합해 '
     '개발 생산성과 품질을 함께 끌어올린 실행 사례이다.',
     bold=True, color=NAVY)

out = 'docs/report/Y_OPS_MONITOR_V2_성과보고서.docx'
doc.save(out)
print('saved', out)
