# -*- coding: utf-8 -*-
"""논의 흐름 정리 발표자료 — 기승전결보다 대화 순서 중심."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt, Emu

HERE = Path(__file__).resolve().parent
OUT = HERE / "Cursor_ABAP_Skills_Parity_발표자료.pptx"
ART = Path("/opt/cursor/artifacts")
ART.mkdir(parents=True, exist_ok=True)

NAVY = RGBColor(0x0B, 0x1F, 0x3A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x5A, 0x5A, 0x5A)
LIGHT = RGBColor(0xF2, 0xF5, 0xF8)


def set_run(p, text, size=18, bold=False, color=DARK):
    p.clear()
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = "Malgun Gothic"


def add_bg(slide, prs, color=LIGHT):
    shape = slide.shapes.add_shape(1, Emu(0), Emu(0), prs.slide_width, prs.slide_height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    spTree = slide.shapes._spTree
    sp = shape._element
    spTree.remove(sp)
    spTree.insert(2, sp)


def add_header_bar(slide, prs):
    bar = slide.shapes.add_shape(1, Emu(0), Emu(0), prs.slide_width, Inches(0.55))
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    t = slide.shapes.add_textbox(Inches(0.55), Inches(0.12), Inches(12), Inches(0.35))
    set_run(t.text_frame.paragraphs[0], "논의 정리  |  Cursor × ABAP Ops Monitor", size=11, color=WHITE)


def add_title(slide, text):
    box = slide.shapes.add_textbox(Inches(0.55), Inches(0.75), Inches(12.2), Inches(0.6))
    set_run(box.text_frame.paragraphs[0], text, size=26, bold=True, color=NAVY)


def add_bullets(slide, lines, top=1.5, size=15, left=0.65):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(12.0), Inches(5.2))
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for line in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(7)
        set_run(p, line, size=size, color=DARK)


def add_table(slide, rows, left, top, width, height, col_w=None):
    cols = len(rows[0])
    table = slide.shapes.add_table(len(rows), cols, Inches(left), Inches(top), Inches(width), Inches(height)).table
    if col_w:
        for i, w in enumerate(col_w):
            table.columns[i].width = Inches(w)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.cell(ri, ci)
            cell.text = str(val)
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(11 if ri else 12)
                    run.font.bold = ri == 0
                    run.font.name = "Malgun Gothic"
                    run.font.color.rgb = WHITE if ri == 0 else DARK
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if ri == 0 else (
                RGBColor(0xE8, 0xEE, 0xF3) if ri % 2 == 0 else WHITE
            )


def add_footer(slide, page, total):
    box = slide.shapes.add_textbox(Inches(0.55), Inches(7.05), Inches(12), Inches(0.3))
    set_run(box.text_frame.paragraphs[0], f"{page} / {total}", size=10, color=MUTED)


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    slides = []

    # 1 표지
    s = prs.slides.add_slide(blank)
    add_bg(s, prs, NAVY)
    b = s.shapes.add_textbox(Inches(0.7), Inches(2.3), Inches(12), Inches(1))
    set_run(b.text_frame.paragraphs[0], "지금까지 논의한 내용 정리", size=32, bold=True, color=WHITE)
    b2 = s.shapes.add_textbox(Inches(0.7), Inches(3.4), Inches(12), Inches(1.2))
    set_run(
        b2.text_frame.paragraphs[0],
        "질문 집계 → A/B 세션 비교 → 채점 방식 → Hard Gate 보정\n"
        "(기승전결보다, 대화가 진행된 순서대로)",
        size=16,
        color=RGBColor(0xB8, 0xC8, 0xD8),
    )
    slides.append(s)

    # 2 흐름
    s = prs.slides.add_slide(blank)
    add_bg(s, prs)
    add_header_bar(s, prs)
    add_title(s, "오늘 정리하는 흐름")
    add_bullets(
        s,
        [
            "① 이 AI 폴더에서 무엇을 물었나 (질문 유형·개수)",
            "② 집계 범위 — 어떤 세션까지 셀 수 있었나",
            "③ Session A와 B의 관계 (Skills를 만들어 재현에 씀)",
            "④ DoD / ‘만족할 때까지 질문 수’가 객관적이지 않다는 문제",
            "⑤ Parity 채점표로 품질을 다시 잼",
            "⑥ 질문 수에 따라 점수가 어떻게 올랐는지",
            "⑦ Hard Gate — 에러·미동작이면 PASS 불가 (곡선 보정)",
            "⑧ 일반화에 대한 추정과 한계",
        ],
        size=16,
    )
    slides.append(s)

    # 3 질문 유형
    s = prs.slides.add_slide(blank)
    add_bg(s, prs)
    add_header_bar(s, prs)
    add_title(s, "① 주로 한 질문 유형")
    add_bullets(
        s,
        [
            "SE38 붙여넣기용 전체/부분 소스 요청 (GitHub·경로 접근 어려움)",
            "컴파일·활성화 오류 수정 (타입, ECC API, DATUM, 문자열 템플릿…)",
            "UX·차트·레이아웃 다듬기 (비율, 색, 툴바, 팝업)",
            "실기 증상 (빈 화면, 팝업 X, ST22 미표시, 더블클릭 동작)",
            "Selection Screen / Dynpro 설정 가이드",
            "후반: 질문 요약·개수 집계·A/B 비교·채점 방식 논의",
        ],
        size=15,
    )
    slides.append(s)

    # 4 개수
    s = prs.slides.add_slide(blank)
    add_bg(s, prs)
    add_header_bar(s, prs)
    add_title(s, "① 질문 개수 (접근 가능 2세션 합산)")
    add_table(
        s,
        [
            ["항목", "개수", "비고"],
            ["총 질문 (메타·시스템 제외)", "약 124~126", "요약/집계 질문 제외"],
            ["오류 관련 (중복 제거)", "23", "같은 오류 재보고는 1건"],
            ["단순 기능 수정 (중복 제거)", "13~15", "UX·컬럼·색·팝업 polish 등"],
        ],
        0.6,
        1.6,
        12.0,
        2.4,
        [4.5, 2.5, 5.0],
    )
    add_bullets(
        s,
        [
            "통짜 소스 / `계속` / 발표자료 / 권한·설계 확정은 ‘총 질문’에는 들어가나",
            "오류·기능수정 버킷에는 넣지 않음",
        ],
        top=4.4,
        size=14,
    )
    slides.append(s)

    # 5 범위
    s = prs.slides.add_slide(blank)
    add_bg(s, prs)
    add_header_bar(s, prs)
    add_title(s, "② 집계 범위 — 보이는 세션 ≠ 전부 열리는 세션")
    add_bullets(
        s,
        [
            "집계에 들어간 것: Awesome skills automation + 통합 운영 모니터링 대시보드",
            "UI에는 약 5개 세션이 보여도, 전사(transcript) API로는 위 2개만 열림",
            "SELECT-OPTIONS 설계서 세션 등(PR #6 등)은 프로그램 관련이지만",
            "  → 만료/접근 불가로 개수에 못 넣음",
            "즉, ‘이 폴더의 모든 과거 대화’가 아니라 ‘지금 읽을 수 있는 대화’ 기준",
        ],
        size=15,
    )
    slides.append(s)

    # 6 A/B 관계
    s = prs.slides.add_slide(blank)
    add_bg(s, prs)
    add_header_bar(s, prs)
    add_title(s, "③ Session A → Skills → Session B")
    add_table(
        s,
        [
            ["", "A Awesome skills", "B 통합 운영 모니터"],
            ["한 일", "장기 구현 + 오류/UX 루프", "설계서+Skills로 Greenfield 재구현"],
            ["끝에", "Preflight·UX·Intake·Delivery 정착", "같은 규칙을 적용해 재현"],
            ["공정 Q", "≈63", "≈16"],
            ["E_compile", "~18+", "2"],
        ],
        0.55,
        1.55,
        12.2,
        3.2,
        [2.2, 5.0, 5.0],
    )
    add_bullets(
        s,
        ["맞음: A에서 겪은 오류·취향을 Skill로 고착하고, B에서 그걸로 다시 만든 구조"],
        top=5.1,
        size=15,
    )
    slides.append(s)

    # 7 Skills 내용
    s = prs.slides.add_slide(blank)
    add_bg(s, prs)
    add_header_bar(s, prs)
    add_title(s, "③ 구체적으로 추가된 것")
    add_bullets(
        s,
        [
            "Rules: 읽기전용 · SE38 붙여넣기 · 선언부+구현 동시",
            "Preflight P1~P10: 선언 누락, ECC fcode, IGS 추측 금지, CSS {}, 타입…",
            "Ops Monitor UX: HTML 차트/팝업, 영역색, ALV 최소 툴바",
            "Intake: 모호한 ‘있어보이게’를 ≤8 선택지로 한 번 계약",
            "B에서 줄어든 것: IGS API 루프, MC_FC_*, LS-DATUM 장루프, MESSAGE 팝업 탐색 등",
            "B에도 남은 것: MANDT/WRITE, 부분 붙여넣기 잔재, Dynpro 수작업, 팝업X, ST22 공백…",
        ],
        size=15,
    )
    slides.append(s)

    # 8 객관성 문제
    s = prs.slides.add_slide(blank)
    add_bg(s, prs)
    add_header_bar(s, prs)
    add_title(s, "④ 객관성 문제")
    add_bullets(
        s,
        [
            "기존: 사용자가 만족하면 중단하고 질문 개수만 셈 → 주관적",
            "DoD도 ‘우리가 만든 내부 체크리스트’이지 외부 공인 지표가 아님",
            "그래서: 품질(Parity)과 과정(질문·오류 수)을 분리하자고 정리",
            "품질 = 동일 설계서 대비 채점표 점수",
            "과정 = 그 점수에 도달하는 데 든 질문·오류 횟수",
        ],
        size=16,
    )
    slides.append(s)

    # 9 Parity 채점
    s = prs.slides.add_slide(blank)
    add_bg(s, prs)
    add_header_bar(s, prs)
    add_title(s, "⑤ Parity 채점표로 다시 잼")
    add_table(
        s,
        [
            ["섹션", "만점"],
            ["S1 기능 (FR)", "20"],
            ["S2 안전·활성화", "12"],
            ["S3 UX 바", "16"],
            ["S4 전달·가이드", "8"],
            ["합계", "56"],
        ],
        0.6,
        1.5,
        6.5,
        3.5,
        [4.0, 2.5],
    )
    add_bullets(
        s,
        [
            "최종 채점: A=54, B=55 (Δ=1)",
            "→ 둘 다 PASS, 품질은 동등",
            "질문 수는 품질 점수에 넣지 않음",
            "동등성: |A−B|≤4 이고 둘 다 PASS",
        ],
        left=7.5,
        top=1.5,
        size=15,
    )
    slides.append(s)

    # 10 곡선
    s = prs.slides.add_slide(blank)
    add_bg(s, prs)
    add_header_bar(s, prs)
    add_title(s, "⑥ 질문 수에 따른 Parity 향상")
    add_bullets(
        s,
        [
            "가로축: 프로그램 관련 누적 질문 (PPT/`계속`/메타 집계 제외)",
            "A: 천천히 오름 → 첫 PASS ≈ Q57 → 최종 54",
            "B: 소스 먼저 나오고 실기에서 막힘 → 첫 PASS ≈ Q12 → 최종 55",
            "같은 Q=16에서: A≈33(FAIL) / B=55(PASS) — 도달 속도 차이",
        ],
        size=16,
    )
    slides.append(s)

    # 11 Hard Gate
    s = prs.slides.add_slide(blank)
    add_bg(s, prs)
    add_header_bar(s, prs)
    add_title(s, "⑦ Hard Gate — 지적하신 대로 보정")
    add_bullets(
        s,
        [
            "문제: B를 ‘처음부터 거의 완성’처럼 채점함 (코드에 HTML 있음 ≈ 고득점)",
            "원칙: 에러 나거나 기능이 안 되면 PASS 안 됨",
            "  · 활성화 실패 → PASS 불가",
            "  · 팝업 X 불능 / ST22 조회 버그 / 빈 화면 → PASS 불가",
            "  · 미실기(코드만) → 점수 상한 40",
            "보정 후 B: Q1=30(FAIL) … Q12=53(첫 PASS) … Q16=55",
            "‘Q=3에 PASS’ 서술은 철회",
        ],
        size=15,
    )
    slides.append(s)

    # 12 일반화
    s = prs.slides.add_slide(blank)
    add_bg(s, prs)
    add_header_bar(s, prs)
    add_title(s, "⑧ 일반화에 대한 추정 — 대체로 맞음")
    add_bullets(
        s,
        [
            "큰 효과의 상당 부분 = 동일 프로그램 유형 오류를 목록화한 덕분",
            "Preflight 약 70%는 다른 ALV/HTML 리포트에도 쓸 만함",
            "UX·도메인 스킬 약 80%+는 Ops Monitor 특화",
            "다른 형식 프로그램이면 오류·질문이 다시 늘 가능성 큼 (아직 실험 n=1)",
            "다양한 프로그램 대화를 쌓아 규칙으로 올리면 효과가 넓어질 것",
        ],
        size=15,
    )
    slides.append(s)

    # 13 한 장 요약
    s = prs.slides.add_slide(blank)
    add_bg(s, prs)
    add_header_bar(s, prs)
    add_title(s, "한 장으로 보면")
    add_table(
        s,
        [
            ["말하고자 하는 것", "숫자/판정"],
            ["최종 품질", "A 54 ≈ B 55 (동등)"],
            ["과정", "공정 Q 63→16, 첫 PASS 57→12"],
            ["방법", "실패 → Skill/Rule 고착 → 재현"],
            ["채점", "Parity 56점 + Hard Gate"],
            ["주의", "코드만으로 PASS 금지 / 유형 1건 한계"],
        ],
        0.6,
        1.55,
        12.1,
        4.2,
        [5.5, 6.6],
    )
    slides.append(s)

    total = len(slides)
    for i, sl in enumerate(slides, 1):
        if i == 1:
            continue
        add_footer(sl, i, total)

    prs.save(OUT)
    dest = ART / OUT.name
    dest.write_bytes(OUT.read_bytes())
    # zip bundle
    import zipfile

    zpath = ART / "Cursor_ABAP_Skills_Parity_bundle.zip"
    with zipfile.ZipFile(zpath, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.write(OUT, OUT.name)
        for name in (
            "Cursor_ABAP_Skills_Parity_handout.md",
            "parity-score-A-vs-B-2026-08.md",
            "parity-score-evolution-vs-Q-2026-08.md",
        ):
            p = HERE / name
            if p.exists():
                zf.write(p, name)
        guide = HERE.parent / "guides" / "ops-monitor-parity-scoring.md"
        if guide.exists():
            zf.write(guide, guide.name)
    print(OUT)
    print(dest)
    print(zpath)
    return OUT


if __name__ == "__main__":
    build()
