#!/usr/bin/env python3
"""세션 분석 보고서 → 4장 발표자료(.pptx) 생성 스크립트.

참조 금지 영역(_private-analysis/)의 산출물이다. 하네스와 무관하다.
실행: python3 make_deck.py
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

FONT = "Malgun Gothic"

INK = RGBColor(0x1F, 0x29, 0x33)       # 본문 먹색
GRAY = RGBColor(0x64, 0x74, 0x8B)      # 보조 설명
LIGHT = RGBColor(0x94, 0xA3, 0xB8)     # 라벨
BLUE = RGBColor(0x13, 0x4E, 0x8F)      # 강조(AI)
BLUE_BG = RGBColor(0xE8, 0xF0, 0xFA)
ORANGE = RGBColor(0xB4, 0x3C, 0x0A)    # 강조(사람 직접)
ORANGE_BG = RGBColor(0xFD, 0xEE, 0xE4)
GREEN = RGBColor(0x14, 0x6B, 0x4B)
PANEL = RGBColor(0xF4, 0xF6, 0xF9)
LINE = RGBColor(0xD9, 0xE0, 0xE8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

W, H = 13.333, 7.5
ML = 0.75                 # 좌우 여백
CW = W - ML * 2           # 본문 폭


def kfont(run, size=14, bold=False, color=INK, name=FONT):
    """한글이 깨지지 않도록 라틴·동아시아·복합 글꼴을 모두 지정한다."""
    f = run.font
    f.name = name
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    anchor = rPr.find(qn("a:latin"))
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            if anchor is not None:
                anchor.addnext(el)
            else:
                rPr.append(el)
        el.set("typeface", name)
        anchor = el


def textbox(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    """lines = [(텍스트, 크기, 굵기, 색, 줄간격배수, 단락앞여백pt), ...]"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, spec in enumerate(lines):
        text, size, bold, color = spec[0], spec[1], spec[2], spec[3]
        spacing = spec[4] if len(spec) > 4 else 1.15
        space_before = spec[5] if len(spec) > 5 else 0
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        if space_before:
            p.space_before = Pt(space_before)
        kfont(p.add_run(), size, bold, color)
        p.runs[0].text = text
    return tb


def rect(slide, x, y, w, h, fill=None, line=None, line_w=1.0, radius=None):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    sh = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius:
        sh.adjustments[0] = radius
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = line
        sh.line.width = Pt(line_w)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    sh.text_frame.text = ""
    return sh


def header(slide, label, title, subtitle):
    textbox(slide, ML, 0.42, CW, 0.3, [(label, 11, True, LIGHT)])
    textbox(slide, ML, 0.74, CW, 0.65, [(title, 31, True, INK)])
    textbox(slide, ML, 1.52, CW, 0.4, [(subtitle, 15, False, GRAY)])
    rect(slide, ML, 2.02, CW, 0.022, fill=LINE)


def footnote(slide, text, y=6.78):
    rect(slide, ML, y - 0.16, CW, 0.018, fill=LINE)
    textbox(slide, ML, y, CW, 0.4, [(text, 11.5, False, GRAY)])


def new_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def notes(slide, text):
    """발표자 노트 (슬라이드 쇼에서는 보이지 않는다)."""
    tf = slide.notes_slide.notes_text_frame
    tf.text = ""
    for i, line in enumerate([l for l in text.strip().split("\n")]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        kfont(p.add_run(), 11, False, INK)
        p.runs[0].text = line


# ---------------------------------------------------------------- 1장
def slide1(prs):
    s = new_slide(prs)
    header(
        s,
        "1 / 4 · 요약",
        "SAP 재고 조회 프로그램을 AI로 만들었습니다",
        "결론부터 — 사람이 8일 걸릴 일을 4시간 30분에 끝냈습니다",
    )

    cards = [
        ("사람이 직접 만들면", "65시간", "약 8일치 일거리   (추정)", ORANGE, ORANGE_BG),
        ("이번 AI 작업", "4시간 30분", "실제 작업 시간   (실측)", BLUE, BLUE_BG),
        ("차이", "약 14배", "빠르게 끝났습니다", GREEN, PANEL),
    ]
    cw, gap = 3.65, 0.44
    for i, (cap, big, note, color, bg) in enumerate(cards):
        x = ML + i * (cw + gap)
        rect(s, x, 2.32, cw, 1.92, fill=bg, radius=0.08)
        rect(s, x, 2.32, 0.055, 1.92, fill=color)
        textbox(s, x + 0.34, 2.56, cw - 0.6, 0.3, [(cap, 13, True, color)])
        textbox(s, x + 0.34, 2.93, cw - 0.6, 0.7, [(big, 40, True, color)])
        textbox(s, x + 0.34, 3.72, cw - 0.6, 0.3, [(note, 12, False, GRAY)])

    facts = [
        ("만든 것", "재고 현황을 단계별로 펼쳐보는 조회 화면 1개 — ZMM_STOCK_TREE01"),
        ("규모", "프로그램 1,479줄 + 설명 문서 약 3,000줄 + 화면 그림 3장"),
        ("진행 기간", "이틀. 다른 업무와 병행했고, 실제 작업 시간만 4시간 30분"),
        ("오류", "SAP에 넣은 뒤 발견된 오류 5건 — 모두 해결 완료"),
    ]
    for i, (k, v) in enumerate(facts):
        y = 4.58 + i * 0.53
        rect(s, ML, y + 0.09, 0.1, 0.1, fill=BLUE)
        textbox(s, ML + 0.28, y, 1.5, 0.3, [(k, 13.5, True, INK)])
        textbox(s, ML + 2.0, y, CW - 2.0, 0.3, [(v, 13.5, False, GRAY)])

    footnote(
        s,
        "AI는 SAP에 접속하지 않습니다. 설계·코드·문서는 AI가 만들고, SAP 입력과 테스트는 사람이 했습니다.",
    )

    notes(s, """
[1장 · 요약]  약 1분

- SAP 재고 현황을 펼쳐보는 조회 화면 1개를 AI와 함께 만들었습니다.
- 결론부터 말씀드리면, 사람이 직접 만들면 약 65시간 걸릴 분량을 실제 작업 4시간 30분에 끝냈습니다.
- 65시간은 추정값이고 4시간 30분은 실측값입니다. 뒤에서 근거를 설명하겠습니다.
- 한 가지 전제를 먼저 말씀드립니다. AI는 SAP에 접속하지 않습니다.
  설계·코드·문서는 AI가 만들었고, SAP에 넣고 돌려보는 일은 전부 사람이 했습니다.
  그래서 '사람을 대체한 수치'가 아니라 '역할을 나눈 결과'로 보셔야 합니다.
""")


# ---------------------------------------------------------------- 2장
def slide2(prs):
    s = new_slide(prs)
    header(
        s,
        "2 / 4 · 진행 방식",
        "묻고 답하며 5단계로 진행했습니다",
        "코드부터 쓰지 않았습니다. 확인된 것만 코드에 넣었습니다",
    )

    steps = [
        ("1", "무엇을 만들지\n묻기", "질문 26개를 주고받았습니다.\n\n모른다고 답한 6개는 추측값으로 메우고 '추정'이라고 표시했습니다."),
        ("2", "화면 그림을\n먼저 확인", "코드를 쓰기 전에 화면 그림을 그려 보여드리고 확인받았습니다.\n\n합계 계산 방식도 3가지 안을 비교해 골랐습니다."),
        ("3", "실제 항목명\n확인", "확인용 작은 프로그램을 먼저 만들어 드렸습니다.\n\n재고 테이블 7개를 직접 조회해 항목명 24개를 확정했습니다."),
        ("4", "전체 코드\n작성", "1,479줄을 한 번에 받아 SAP에 붙여넣었습니다.\n\n주석은 한국어와 영어를 함께 달았습니다."),
        ("5", "오류 고치고\n화면 만들기", "9번 주고받으며 오류 5건을 고치고, SAP 화면·버튼·라벨 정의를 받았습니다."),
    ]
    gap = 0.3
    cw = (CW - gap * (len(steps) - 1)) / len(steps)
    for i, (num, title, body) in enumerate(steps):
        x = ML + i * (cw + gap)
        rect(s, x, 2.34, cw, 2.72, fill=PANEL, line=LINE, radius=0.05)
        rect(s, x, 2.34, cw, 0.055, fill=BLUE)
        textbox(s, x + 0.22, 2.54, 0.4, 0.34, [(num, 19, True, BLUE)])
        textbox(s, x + 0.22, 2.94, cw - 0.44, 0.72,
                [(ln, 14.5, True, INK) for ln in title.split("\n")])
        textbox(s, x + 0.22, 3.72, cw - 0.44, 1.2,
                [(ln, 11.5, False, GRAY, 1.25) for ln in body.split("\n")])

    rect(s, ML, 5.32, CW, 1.18, fill=BLUE_BG, radius=0.05)
    rect(s, ML, 5.32, 0.055, 1.18, fill=BLUE)
    textbox(s, ML + 0.32, 5.52, CW - 0.64, 0.3,
            [("핵심 — 모르는 것을 추측해서 코드에 넣지 않았습니다", 15, True, BLUE)])
    textbox(s, ML + 0.32, 5.88, CW - 0.64, 0.5,
            [("확인용 프로그램을 먼저 돌린 덕분에, 추측으로는 알 수 없는 4가지를 미리 걸러냈습니다 — "
              "없는 항목 2건, 수량이 두 번 더해질 위험 1건, 금액 계산식 1건.", 12.5, False, INK, 1.3)])

    footnote(s, "1~3단계는 첫째 날, 4~5단계는 둘째 날에 진행했습니다.")

    notes(s, """
[2장 · 진행 방식]  약 1분 30초

- 다섯 단계로 진행했습니다. 핵심은 '코드부터 쓰지 않았다'는 점입니다.
- 1단계, 무엇을 만들지 질문 26개로 확인했습니다. 모른다고 답한 6개는 추측값으로 채운 뒤
  '추정'이라고 표시해 두고, 코드에는 넣지 않았습니다.
- 2단계, 코드를 쓰기 전에 화면 그림을 먼저 그려 확인받았습니다. 합계 계산 방식은 3가지 안을 비교해 골랐습니다.
- 3단계가 가장 중요합니다. 재고 항목의 실제 이름을 확인하려고 '확인용 작은 프로그램'을 먼저 만들어 드렸습니다.
  그걸 돌려본 결과, 추측으로는 알 수 없는 4가지가 나왔습니다.
  특히 수량이 두 번 더해질 위험이 하나 있었는데, 이건 오류가 안 나고 숫자만 틀리는 종류입니다.
- 4단계에서 1,479줄을 한 번에 받아 붙여넣었고, 5단계에서 오류를 고치고 화면을 만들었습니다.
""")


# ---------------------------------------------------------------- 3장
def slide3(prs):
    s = new_slide(prs)
    header(
        s,
        "3 / 4 · 수치",
        "숫자로 보면",
        "주고받은 43번 중 72%는 제대로 진행된 대화였습니다",
    )

    # 좌측 — 대화 구성
    lx = ML
    textbox(s, lx, 2.28, 5.3, 0.34, [("주고받은 횟수  43번", 17, True, INK)])
    bars = [
        ("제대로 진행", 31, "72%", BLUE),
        ("오류 고치기", 5, "12%", ORANGE),
        ("AI가 빠뜨려 다시 요청", 4, "9%", ORANGE),
        ("헛턴 (다시 출력 요청 등)", 3, "7%", LIGHT),
    ]
    max_w = 2.1
    for i, (name, val, pct, color) in enumerate(bars):
        y = 2.82 + i * 0.63
        textbox(s, lx, y, 2.4, 0.3, [(name, 12.5, False, INK)])
        bw = max(0.12, max_w * val / 31)
        rect(s, lx + 2.45, y + 0.025, max_w, 0.24, fill=RGBColor(0xED, 0xF1, 0xF5))
        rect(s, lx + 2.45, y + 0.025, bw, 0.24, fill=color)
        textbox(s, lx + 4.65, y, 0.55, 0.3, [(f"{val}번", 12.5, True, color)], align=PP_ALIGN.RIGHT)
        textbox(s, lx + 2.45, y + 0.3, max_w, 0.26, [(pct, 10.5, False, GRAY)])

    rect(s, lx, 5.46, 5.2, 1.0, fill=PANEL, radius=0.07)
    textbox(s, lx + 0.3, 5.66, 4.6, 0.3, [("오류 5건이 나온 곳", 13, True, INK)])
    textbox(s, lx + 0.3, 6.0, 4.6, 0.6,
            [("5건 중 3건이 '화면에 표를 펼쳐 보여주는 기능'의", 11.5, False, GRAY, 1.3),
             ("사용법을 알아내는 과정에서 나왔습니다.", 11.5, False, GRAY, 1.3)])

    # 우측 — 만든 것
    rx = ML + 5.65
    rw = CW - 5.65
    textbox(s, rx, 2.28, rw, 0.34, [("만든 것 · 다룬 범위", 17, True, INK)])
    rows = [
        ("프로그램", "1,479줄", "이 중 주석이 25%"),
        ("설명 문서", "약 3,000줄", "화면·항목 정의 포함"),
        ("화면 그림", "3장", "미리 보여드린 것"),
        ("조회 테이블", "11개", "항목 24개 직접 확인"),
        ("실행 결과", "재고 69만 건", "화면에는 7,339줄"),
        ("저장한 교훈", "7건", "다음 작업에 자동 반영"),
    ]
    for i, (k, v, note) in enumerate(rows):
        y = 2.78 + i * 0.63
        if i % 2 == 0:
            rect(s, rx, y - 0.08, rw, 0.6, fill=PANEL)
        textbox(s, rx + 0.2, y + 0.03, 1.55, 0.3, [(k, 12.5, False, GRAY)])
        textbox(s, rx + 1.85, y, 1.7, 0.34, [(v, 14.5, True, BLUE)])
        textbox(s, rx + 3.65, y + 0.03, rw - 3.85, 0.3, [(note, 11.5, False, GRAY)])

    footnote(
        s,
        "되돌아간 12번(28%) 중 4번은 AI가 'SAP 화면에서 손으로 만들 부분'의 입력값을 표로 주지 않아 생긴 것입니다. 규칙에 추가해 막았습니다.",
    )

    notes(s, """
[3장 · 수치]  약 1분 30초

- 총 43번 주고받았습니다. 이 중 31번, 72%는 제대로 진행된 대화입니다.
- 오류 고치기 5번. 이 중 3건이 '화면에 표를 펼쳐 보여주는 기능'의 사용법을 알아내는 과정에서 나왔습니다.
  처음 쓰는 기능이라 생긴 비용입니다.
- AI가 빠뜨려서 다시 요청한 것이 4번 있었습니다. 원인은 하나였습니다. 뒤에서 설명하겠습니다.
- 헛턴 3번은 같은 내용을 다시 출력해 달라고 한 것으로, 진행에는 도움이 안 된 횟수입니다.
- 오른쪽은 만든 것입니다. 실제 실행했을 때 재고 69만 건을 읽어 화면에는 7,339줄로 요약해 보여줍니다.
- 마지막 줄, 교훈 7건을 저장했습니다. 다음 프로그램을 만들 때 AI가 자동으로 참고합니다.
""")


# ---------------------------------------------------------------- 4장
def slide4(prs):
    s = new_slide(prs)
    header(
        s,
        "4 / 4 · 결론",
        "시간 비교와 배운 점",
        "선행 투자가 끝났으므로, 다음 프로그램은 더 빨라집니다",
    )

    # 좌측 — 시간 비교
    lx, lw = ML, 5.75
    textbox(s, lx, 2.28, lw, 0.34, [("같은 프로그램을 만드는 데 걸리는 시간", 16, True, INK)])
    comp = [
        ("사람이 직접 (AI 없이)", 65.0, "65시간", ORANGE, "추정 · 범위 40~100시간"),
        ("이번 AI 작업", 4.6, "4시간 30분", BLUE, "실측 · 이틀에 걸쳐 병행"),
        ("같은 것을 다시 만들면", 3.0, "3시간", GREEN, "추정 · 교훈 7건이 이미 반영됨"),
    ]
    bar_max = 4.0
    for i, (name, val, label, color, note) in enumerate(comp):
        y = 2.86 + i * 1.22
        textbox(s, lx, y, lw, 0.3, [(name, 13, True, INK)])
        bw = max(0.1, bar_max * val / 65.0)
        rect(s, lx, y + 0.36, bar_max, 0.3, fill=RGBColor(0xED, 0xF1, 0xF5))
        rect(s, lx, y + 0.36, bw, 0.3, fill=color)
        textbox(s, lx + 4.12, y + 0.34, 1.6, 0.34, [(label, 15, True, color)])
        textbox(s, lx, y + 0.74, lw, 0.3, [(note, 11, False, GRAY)])

    # 우측 — 배운 점
    rx = ML + 6.35
    rw = CW - 6.35
    textbox(s, rx, 2.28, rw, 0.34, [("배운 점", 16, True, INK)])
    lessons = [
        ("모르는 것은 먼저 확인했습니다",
         "가장 위험한 것은 오류가 아니라 '오류 없이 조용히 틀린 숫자'입니다. 확인용 프로그램 1개로 막았습니다."),
        ("되돌아간 4번은 원인이 하나였습니다",
         "SAP 화면에서 손으로 만드는 부분의 입력값을 표로 주지 않았습니다. 규칙에 넣어 다시 생기지 않게 했습니다."),
        ("세 번째 프로그램부터 본격적인 이득입니다",
         "규칙과 교훈을 쌓는 초기 투자가 끝났습니다. 반복할수록 시간 차이가 커집니다."),
    ]
    for i, (title, body) in enumerate(lessons):
        y = 2.82 + i * 1.3
        rect(s, rx, y, rw, 1.14, fill=PANEL, radius=0.07)
        rect(s, rx, y, 0.055, 1.14, fill=BLUE)
        textbox(s, rx + 0.42, y + 0.17, 0.3, 0.3, [(str(i + 1), 14, True, BLUE)])
        textbox(s, rx + 0.78, y + 0.17, rw - 1.0, 0.3, [(title, 13.5, True, INK)])
        textbox(s, rx + 0.78, y + 0.53, rw - 1.0, 0.55, [(body, 11.5, False, GRAY, 1.3)])

    footnote(
        s,
        "주의 — 65시간은 통상 생산성 기준의 추정값입니다. 사내 유사 개발 실적이 있으면 그 값으로 바꿔야 정확합니다. "
        "대기·협의·승인·이송 시간은 포함하지 않았습니다.",
    )

    notes(s, """
[4장 · 결론]  약 1분 30초

- 시간 비교입니다. 사람이 직접 65시간, 이번 AI 작업 4시간 30분, 같은 것을 다시 만들면 3시간으로 봅니다.
- 65시간 산정 근거는 요구사항 확정 6시간, 기존 화면 분석 6시간, 항목 조사 6시간,
  설계 4시간, 코딩 16~24시간, 화면 기능 8시간, 테스트 8시간, 문서 4시간입니다.
- 배운 점 세 가지입니다.
  첫째, 모르는 것은 추측하지 않고 먼저 확인했습니다. 가장 위험한 건 오류가 아니라
  오류 없이 조용히 틀린 숫자입니다. 확인용 프로그램 하나로 막았습니다.
  둘째, 되돌아간 4번은 원인이 하나였습니다. SAP 화면에서 사람이 손으로 만들어야 하는 부분의
  입력값을 표로 주지 않았던 것입니다. 규칙에 넣어 다시 생기지 않게 했습니다.
  셋째, 세 번째 프로그램부터 본격적인 이득입니다. 규칙을 쌓는 초기 투자가 이미 끝났습니다.
- 마지막으로 주의사항입니다. 65시간은 통상 생산성 기준의 추정값입니다.
  사내 유사 개발 실적이 있으면 그 값으로 바꿔야 정확합니다.
  대기·협의·승인·이송 시간은 포함하지 않은 순작업 기준입니다.
""")

def main():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    for fn in (slide1, slide2, slide3, slide4):
        fn(prs)
    out = "20261001-session-deck.pptx"
    prs.save(out)
    print(f"생성 완료: {out} ({len(prs.slides.__iter__.__self__._sldIdLst)}장)")


if __name__ == "__main__":
    main()
